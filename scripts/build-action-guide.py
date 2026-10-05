"""
Build the "What to do" data from the Risk Action Guide Book (data-source/RISK_ACTION_GUIDE_BOOK_0222.docx).

  python scripts/build-action-guide.py

Writes, all bilingual (English and Kiswahili, as published in the guide), loaded by the site on demand:
  src/data/action-guide/hazards/<hazard>.json one file per hazard: its three alert levels (impact -> actions)
  src/data/action-guide/incidents.json        the incident and accident rapid-response guide
  src/data/action-guide/regions/<region>.json per council (by site council id): the hazards the guide
                                              documents for it and its "Know your risk" statements;
                                              Tabora's councils also carry their own alert-level actions

The per-council alert tables of the guide are only used to read each council's documented hazards: their
generated sentences repeat one flood template (and some substitutions are broken), so the site shows the
national hazard actions instead, except for Tabora, whose councils have hand-written tables.
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'data-source' / 'RISK_ACTION_GUIDE_BOOK_0222.docx'
OUT = ROOT / 'src' / 'data' / 'action-guide'
COUNCILS = json.loads((ROOT / 'src' / 'data' / 'tanzania-councils-index.json').read_text(encoding='utf-8'))

HAZARDS = {
    'HEAVY RAINFALL': 'heavyRainfall',
    'FLOODS': 'floods',
    'LANDSLIDE': 'landslide',
    'STRONG WINDS': 'strongWinds',
    'LARGE WAVES': 'largeWaves',
    'WILDFIRE': 'wildfire',
    'DROUGHT': 'drought',
    'EARTHQUAKE AND TSUNAMI': 'earthquake',
    'PUBLIC HEALTH': 'publicHealth',
}

# Words in a council's documented hazard line -> guide hazard. Order matters only for readability.
HAZARD_WORDS = [
    ('floods', r'flood'),
    ('heavyRainfall', r'heavy rain|rainfall|downpour|\brains?\b|lightning|thunder|storm ?water|rainstorm|heavy storms?'),
    ('landslide', r'landslide|mudflow|rockfall|rock fall|slope failure|gull(y|ies)'),
    ('strongWinds', r'strong wind|\bwinds?\b|gusts?\b|strong storms?|windstorm|cyclone|whirlwind'),
    ('largeWaves', r'wave|surge|storm tide|boat|ferry|ferries|coastal|sea-level|sea level|shore erosion|coastal erosion|lake storms'),
    ('wildfire', r'wildfire|bush ?fires?|forest fires?|grass ?fires?|veld ?fires?|(dry[- ]season|plantation|miombo|rangeland|range) fires?|fires? spreads?'),
    ('drought', r'drought|dry spell|water shortage|crop failure|food shortage|famine'),
    ('earthquake', r'earthquake|tremor|tsunami|seismic'),
    ('publicHealth', r'cholera|disease|epidemic|outbreak|malaria|dengue|typhoid|mpox|marburg|ebola|anthrax|rabies|plague|diarrho'),
]

# Guide councils listed under an older name, or merged into one census council (normalised guide name ->
# site council name). Several guide councils may point at one site council: their statements combine.
ALIASES = {
    'ilalamunicipalcouncil': 'Dar Es Salaam City',  # Ilala Municipal Council became Dar es Salaam City Council (2021)
    'zanzibarmunicipalcouncil': 'Mjini Municipal',  # Stone Town and the Zanzibar urban council are the Mjini council
    'kilomberodistrictcouncil': 'Mlimba District',  # Kilombero District Council was renamed Mlimba
    'mpandadistrictcouncil': 'Tanganyika District',  # Mpanda District Council was renamed Tanganyika
    'mtwaramunicipalcouncil': 'Mtwara Mikindani Municipal',
    'lindidistrictcouncil': 'Mtama District',  # the former Lindi District Council area is Mtama District
}


def clean(s: str) -> str:
    """Collapse whitespace and keep to the site's punctuation (no em dashes)."""
    s = re.sub(r'\s+', ' ', s or '').strip()
    s = re.sub(r'\s*\u2014\s*', ', ', s)
    return s


def place_key(name: str) -> str:
    return re.sub(r'[^a-z0-9]', '', (name or '').lower())


SW_TYPE = {'wilaya': 'district', 'mji': 'town', 'manispaa': 'municipal', 'jiji': 'city'}
EN_TYPE = {'district': 'district', 'dc': 'district', 'town': 'town', 'tc': 'town', 'municipal': 'municipal', 'municipality': 'municipal', 'mc': 'municipal', 'city': 'city', 'cc': 'city'}


def parse_council(name: str):
    """'Arusha City Council' -> ('arusha', 'city'); 'Halmashuari Ya Wilaya Ya Mbogwe District' -> ('mbogwe', 'district')."""
    ws = [w for w in re.split(r'[^a-z0-9]+', (name or '').lower()) if w]
    kind = None
    if any(w in ('halmashauri', 'halmashuari') or w in SW_TYPE for w in ws):
        for w in ws:
            if w in SW_TYPE and kind is None:
                kind = SW_TYPE[w]
        ws = [w for w in ws if w not in ('halmashauri', 'halmashuari', 'ya', 'wa', 'la') and w not in SW_TYPE]
    while len(ws) > 1 and ws[-1] == 'council':
        ws = ws[:-1]
    if len(ws) > 1 and ws[-1] in EN_TYPE:
        kind = kind or EN_TYPE[ws[-1]]
        ws = ws[:-1]
    return ''.join(ws), kind


BY_REGION = {}
for c in COUNCILS:
    base, kind = parse_council(c['name'])
    BY_REGION.setdefault(place_key(c['reg']), []).append({'id': c['code'], 'name': c['name'], 'base': base, 'kind': kind})


def assign_councils(region: str, names):
    """Guide council names -> site councils within one region. Renamed councils and exact name-and-type
    matches first; then a name alone, but only against site councils no other guide council has taken
    (so "Mtwara Municipal" can never land on "Mtwara District")."""
    pool = BY_REGION.get(region, [])
    result, taken = {}, set()
    for name in names:
        alias = ALIASES.get(place_key(name))
        if alias:
            hit = next((c for c in pool if c['name'] == alias), None)
            if hit:
                result[name] = hit
            continue
        base, kind = parse_council(name)
        exact = [c for c in pool if c['base'] == base and c['kind'] == kind and c['id'] not in taken]
        if len(exact) == 1:
            result[name] = exact[0]
            taken.add(exact[0]['id'])
    for name in names:
        if name in result or place_key(name) in ALIASES:
            continue
        base, _ = parse_council(name)
        loose = [c for c in pool if c['base'] == base and c['id'] not in taken]
        if len(loose) == 1:
            result[name] = loose[0]
            taken.add(loose[0]['id'])
    return result


def hazards_in(text: str):
    found = []
    low = text.lower()
    for key, rx in HAZARD_WORDS:
        m = re.search(rx, low)
        if m:
            found.append((m.start(), key))
    return [k for _, k in sorted(found)]


# ------------------------------------------------------------------------------------------------ #
# Read the document as blocks: paragraphs and tables (merged cells de-duplicated).
# ------------------------------------------------------------------------------------------------ #

doc = Document(str(SRC))
blocks = []
for child in doc.element.body.iterchildren():
    tag = child.tag.rsplit('}', 1)[-1]
    if tag == 'p':
        text = Paragraph(child, doc).text.strip()
        if text:
            blocks.append({'t': 'p', 'text': text})
    elif tag == 'tbl':
        rows = []
        for row in Table(child, doc).rows:
            seen, cells = set(), []
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                cells.append([q.text.strip() for q in cell.paragraphs if q.text.strip()])
            rows.append(cells)
        blocks.append({'t': 'table', 'rows': rows})


def head(table):
    return ' '.join(' '.join(c) for c in table['rows'][0]).upper()


def level_of(table):
    m = re.search(r'LEVEL\s*([123])', head(table))
    return int(m.group(1)) if m else None


def impact_rows(table):
    """IMPACT (EN) | ATHARI (SW) | ACTION (EN) | HATUA (SW) rows -> [{impact, actions}]."""
    out = []
    for r in table['rows'][1:]:
        if len(r) < 4:
            continue
        impact_en, impact_sw = clean(' '.join(r[0])), clean(' '.join(r[1]))
        actions_en, actions_sw = [clean(x) for x in r[2] if clean(x)], [clean(x) for x in r[3] if clean(x)]
        if not impact_en and not actions_en:
            continue
        out.append({'impact': {'en': impact_en, 'sw': impact_sw}, 'actions': {'en': actions_en, 'sw': actions_sw}})
    return out


def levels_after(start: int, stop: int):
    """Alert-level header tables followed by their impact/action tables, between two block indexes."""
    levels, current = {}, None
    for b in blocks[start:stop]:
        if b['t'] != 'table':
            continue
        if len(b['rows']) == 1 and level_of(b):
            current = level_of(b)
            continue
        if current and 'IMPACT' in head(b):
            levels.setdefault(str(current), []).extend(impact_rows(b))
            current = None
    return levels


def index_of(pred, start=0):
    for i in range(start, len(blocks)):
        if pred(blocks[i]):
            return i
    return len(blocks)


is_p = lambda rx: (lambda b: b['t'] == 'p' and re.match(rx, b['text']))

# ------------------------------------------------------------------------------------------------ #
# The nine hazards and the incident guide.
# ------------------------------------------------------------------------------------------------ #

guide = {
    'source': {
        'title': 'Risk Action Guide Book',
        'basis': 'National Disaster Management Strategy 2022–2027; National Disaster Preparedness and Response Plan 2022',
        'file': 'data-source/RISK_ACTION_GUIDE_BOOK_0222.docx',
    },
    'hazards': {},
    'incidents': [],
}

for title, key in HAZARDS.items():
    start = index_of(is_p(r'^\d+\.0 ' + re.escape(title) + r'$'))
    sectors = index_of(is_p(r'^\d+\.1 SECTOR IMPACTS'), start)
    levels = levels_after(start, sectors)
    assert set(levels) == {'1', '2', '3'}, f'{title}: levels {sorted(levels)}'
    guide['hazards'][key] = {'levels': levels}

inc_start = index_of(is_p(r'^10\.0 '))
inc_stop = index_of(is_p(r'^11\.0 '))
group = None
for b in blocks[inc_start:inc_stop]:
    if b['t'] != 'table':
        continue
    if len(b['rows']) == 1 and len(b['rows'][0]) == 1:
        en, _, sw = clean(' '.join(b['rows'][0][0])).partition('|')
        group = {'name': {'en': en.strip().capitalize(), 'sw': sw.strip().capitalize()}, 'rows': []}
        guide['incidents'].append(group)
    elif group and 'INCIDENT' in head(b):
        for r in b['rows'][1:]:
            group['rows'].append({
                'incident': {'en': clean(' '.join(r[0])), 'sw': clean(' '.join(r[1]))},
                'actions': {'en': [clean(x) for x in r[2] if clean(x)], 'sw': [clean(x) for x in r[3] if clean(x)]},
            })

# ------------------------------------------------------------------------------------------------ #
# Regions: "Know your risk" statements and each council's documented hazards.
# ------------------------------------------------------------------------------------------------ #

region_rx = r'^(\d+)\. (.+?) REGION \| MKOA WA (.+)$'
region_starts = [i for i, b in enumerate(blocks) if b['t'] == 'p' and re.match(region_rx, b['text'])]
COUNCIL_NAME = re.compile(r'(Council|District|Town|Municipal|City)$')

problems = []
regions = {}


def know_rows(table):
    out = []
    for r in table['rows'][1:]:
        if len(r) < 4:
            continue
        en, sw = [clean(x) for x in r[2] if clean(x)], [clean(x) for x in r[3] if clean(x)]
        out.append({
            'area': {'en': clean(' '.join(r[0])), 'sw': clean(' '.join(r[1]))},
            'risk': {'en': en[0] if en else '', 'sw': sw[0] if sw else ''},
            'advice': {'en': ' '.join(en[1:]), 'sw': ' '.join(sw[1:])},
        })
    return out


for n, start in enumerate(region_starts):
    stop = region_starts[n + 1] if n + 1 < len(region_starts) else len(blocks)
    m = re.match(region_rx, blocks[start]['text'])
    region_en = m.group(2).title()
    region = place_key(m.group(2))
    if region not in BY_REGION:
        problems.append(f'region not on the site: {region_en}')
        continue
    entry = regions.setdefault(region, {'region': {'en': f'{region_en} Region', 'sw': clean(f'Mkoa wa {m.group(3).title()}')}, 'councils': {}})
    sub_starts = [i for i in range(start, stop) if blocks[i]['t'] == 'p' and re.match(r'^\d+\.\d+ .+ \| ', blocks[i]['text'])]

    if sub_starts:
        # Tabora: a subsection per council with its own "Know your risk" table and alert levels.
        names = [re.sub(r'^\d+\.\d+ ', '', blocks[s]['text']).split(' | ')[0].title() for s in sub_starts]
        assigned = assign_councils(region, names)
        for k, s in enumerate(sub_starts):
            e = sub_starts[k + 1] if k + 1 < len(sub_starts) else stop
            name_en = names[k]
            council = assigned.get(name_en)
            if not council:
                problems.append(f'{region_en}: no site council for "{name_en}"')
                continue
            know = next((know_rows(b) for b in blocks[s:e] if b['t'] == 'table' and head(b).startswith('AREA')), [])
            levels = levels_after(s, e)
            first = ' '.join(r['impact']['en'] for r in levels.get('1', []))
            entry['councils'][council['id']] = {
                'name': name_en,
                'hazards': hazards_in(first + ' ' + ' '.join(x['risk']['en'] for x in know)),
                'know': know,
                'levels': levels,
            }
        # The regional profile table: setting, priority hazards and who is most at risk, per council.
        profile = next((b for b in blocks[start:sub_starts[0]] if b['t'] == 'table' and 'PRIORITY HAZARDS' in head(b)), None)
        if profile:
            row_names = [' '.join(r[0]) for r in profile['rows'][1:]]
            by_name = assign_councils(region, row_names)
            for r, row_name in zip(profile['rows'][1:], row_names):
                council = by_name.get(row_name)
                if council and council['id'] in entry['councils']:
                    entry['councils'][council['id']]['profile'] = {
                        'setting': clean(' '.join(r[1])),
                        'hazards': clean(' '.join(r[2])),
                        'atRisk': clean(' '.join(r[3])),
                    }
        continue

    tables = [b for b in blocks[start:stop] if b['t'] == 'table']
    know_table = next((t for t in tables if head(t).startswith('AREA')), None)
    know = know_rows(know_table) if know_table else []

    # Council alert tables: a council's first row names it (last paragraph of the first cell).
    order, hazard_lines = [], {}
    for t in (t for t in tables if head(t).startswith('COUNCIL')):
        current = None
        for r in t['rows'][1:]:
            c0 = r[0] if r else []
            if c0 and COUNCIL_NAME.search(c0[-1]):
                current = c0[-1]
                order.append(current)
            level = re.search(r'Level\s*([123])', ' '.join(r[1]) if len(r) > 1 else '')
            if current and level and level.group(1) == '1' and len(r) > 2:
                hazard_lines[current] = clean(' '.join(r[2]))

    assigned = assign_councils(region, order)
    for i, name in enumerate(order):
        council = assigned.get(name)
        if not council:
            problems.append(f'{region_en}: no site council for "{name}"')
            continue
        line = hazard_lines.get(name, '')
        documented = re.split(r'(?<=[a-z\)])\. ', line, maxsplit=1)[0]
        statements = []
        if i < len(know):
            row = know[i]
            prefix = row['area']['en'].split(':', 1)[0]
            p_base, _ = parse_council(prefix)
            n_base, _ = parse_council(name)
            if ':' in row['area']['en'] and p_base != n_base:
                problems.append(f'{region_en}: "Know your risk" row {i + 1} ("{prefix}") paired with "{name}"')
            statements.append(row)
        existing = entry['councils'].get(council['id'])
        if existing:  # two guide councils on one site council (see ALIASES)
            existing['know'].extend(statements)
            existing['hazards'] = list(dict.fromkeys(existing['hazards'] + hazards_in(documented) + [h for row in statements for h in hazards_in(row['risk']['en'])]))
            continue
        # The documented hazard line first, then any hazard the area's own risk statement names.
        found = hazards_in(documented) + [h for row in statements for h in hazards_in(row['risk']['en'])]
        entry['councils'][council['id']] = {'name': name, 'hazards': list(dict.fromkeys(found)), 'documented': documented, 'know': statements}
    if len(know) > len(order):
        problems.append(f'{region_en}: {len(know) - len(order)} "Know your risk" rows without a council')

# ------------------------------------------------------------------------------------------------ #
# Checks and output
# ------------------------------------------------------------------------------------------------ #

covered = {cid for r in regions.values() for cid in r['councils']}
missing = [c for c in COUNCILS if c['code'] not in covered]
for c in missing:
    problems.append(f'site council without guide content: {c["code"]} {c["name"]} ({c["reg"]})')
for r in regions.values():
    for cid, c in r['councils'].items():
        if not c['hazards']:
            problems.append(f'{cid} {c["name"]}: no recognised hazard in "{c.get("documented", "")[:90]}"')


def walk(o):
    if isinstance(o, dict):
        for v in o.values():
            yield from walk(v)
    elif isinstance(o, list):
        for v in o:
            yield from walk(v)
    elif isinstance(o, str):
        yield o


dashes = [s for s in walk([guide, regions]) if '\u2014' in s]
assert not dashes, dashes[:3]

counts = Counter(h for r in regions.values() for c in r['councils'].values() for h in c['hazards'])
print(f'hazards: {len(guide["hazards"])} | incident groups: {len(guide["incidents"])} | regions: {len(regions)} | councils: {len(covered)} of {len(COUNCILS)}')
print('councils per documented hazard:', dict(counts.most_common()))
print(f'problems ({len(problems)}):')
for p in problems:
    print('  -', p)
if problems:
    # Nothing is written until every site council has its guidance (see ALIASES for renamed councils).
    sys.exit(1)


def write(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')


for key, hazard in guide['hazards'].items():
    write(OUT / 'hazards' / f'{key}.json', hazard)
write(OUT / 'incidents.json', guide['incidents'])
for key, r in regions.items():
    write(OUT / 'regions' / f'{key}.json', r)
print(f'written to {OUT.relative_to(ROOT)}')
