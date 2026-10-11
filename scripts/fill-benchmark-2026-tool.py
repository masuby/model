#!/usr/bin/env python3
"""
fill-benchmark-2026-tool.py - the collection tool as the institutions should continue from: the workbook's
raw values (the plain Benchmark) overlaid with the Benchmark 2026 institutional values
(src/data/tanzania-benchmark-2026.json), each in a green cell, plus a 'Benchmark 2026' sheet listing every
dataset with its source, period, resolution, method and institution.

Output: inform_sheets/INFORM_TZ_Collection_Benchmark_2026.xlsx (git-ignored like the other tools; copy it
to inform_sheets/inform_tools/ with the rest).
"""
import json
import os
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BM = json.load(open(os.path.join(ROOT, 'src/data/tanzania-benchmark-2026.json'), encoding='utf-8'))
SPEC = json.load(open(os.path.join(ROOT, 'src/data/inform-indicator-spec.json'), encoding='utf-8'))
for _id, _a in BM['activations'].items():
    SPEC[_id] = {**SPEC[_id], **_a}
name2id = {s['name']: s['id'] for s in SPEC.values() if s.get('use') == 'Yes'}
index = json.load(open(os.path.join(ROOT, 'src/data/tanzania-councils-index.json'), encoding='utf-8'))
region_id = {}
for c in index:
    region_id[c['code']] = 'R-' + ''.join(ch for ch in c['reg'].lower() if ch.isalnum())

src = os.path.join(ROOT, 'inform_sheets/INFORM_TZ_Collection_Benchmark.xlsx')
wb = openpyxl.load_workbook(src)
entry = wb['Council Entry']
HEADR = 10
col_id = {col: name2id[entry.cell(1, col).value] for col in range(4, entry.max_column + 1) if entry.cell(1, col).value in name2id}
by_unit = {}
for key, vals in BM['values'].items():
    d = BM['datasets'][key]
    for unit_id, value in vals.items():
        by_unit.setdefault(unit_id, {})[d['indicator']] = {'value': value}
GREEN = PatternFill('solid', fgColor='E2F0D9')
written = 0
for r in range(HEADR + 1, entry.max_row + 1):
    code = entry.cell(r, 3).value
    if not code:
        continue
    for col, sid in col_id.items():
        v = by_unit.get(code, {}).get(sid) or by_unit.get(region_id.get(code, ''), {}).get(sid)
        if v is None:
            continue
        cell = entry.cell(r, col, v['value'])
        cell.fill = GREEN
        written += 1
# the datasets sheet
ws = wb.create_sheet('Benchmark 2026')
ws.sheet_view.showGridLines = False
ws.append(['INFORM Tanzania - Benchmark 2026 (release %s): institutional values prepared from public datasets' % BM['release']])
ws['A1'].font = Font(bold=True, size=12)
ws.append([BM['note']])
ws['A2'].alignment = Alignment(wrap_text=True)
ws.append([])
hdr = ['Indicator id', 'Indicator', 'Unit', 'Institution', 'Values', 'Level', 'Period', 'Resolution', 'Source', 'Dataset', 'Method', 'Note']
ws.append(hdr)
for c in range(1, len(hdr) + 1):
    ws.cell(4, c).font = Font(bold=True)
    ws.cell(4, c).fill = PatternFill('solid', fgColor='DDEBF7')
for key, d in BM['datasets'].items():
    ws.append([d['indicator'], d.get('name') or SPEC.get(d['indicator'], {}).get('name'), d['unit'], d['institution'], d['values'], d['level'], d['period'], d['resolution'], d['source'], d['dataset'], d['method'], d['note']])
for col, w in zip('ABCDEFGHIJKL', (18, 34, 14, 12, 8, 9, 18, 22, 48, 48, 60, 60)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
ws.append([])
ws.append(['Activated template indicators (Tanzanian definition, range resolved from the data as the workbook does)'])
ws.cell(ws.max_row, 1).font = Font(bold=True)
ws.append(['Indicator id', 'Name', 'Unit', 'Keyed at', 'Direction', 'Normalisation', 'Ref min', 'Ref max', 'Basis', 'Source'])
for sid, a in BM['activations'].items():
    ws.append([sid, a['name'], a['unit'], a['keyed_at'], a['sign'], a['normalisation'], a['resolved_min'], a['resolved_max'], a['basis'], a['source']])
ws.append([])
ws.append(['Green cells on Council Entry are Benchmark 2026 values; white cells are the workbook values; a blank cell has no value yet. Built %s.' % datetime.now(timezone.utc).isoformat(timespec='seconds')])
out = os.path.join(ROOT, 'inform_sheets/INFORM_TZ_Collection_Benchmark_2026.xlsx')
wb.save(out)
print('wrote', out, '| benchmark cells written:', written, '| indicator columns:', len(col_id))
